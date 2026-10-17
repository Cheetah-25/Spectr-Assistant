#!/usr/bin/env python3
"""Build the page-tagged knowledge base from the AFS materials.

Reads the source materials in ``materials/`` and writes page-tagged knowledge
files into ``knowledge/``:

  * S01, S03, S04, S05 : one entry per PDF page (1-based; title page = 1), with a
    scope tag, written to ``knowledge/<Sxx>.json``.
  * S02                : loaded from the committed ``knowledge/S02_slides.json``
    (slide notes; ``S02 p.N`` = slide N). No vision pass.
  * xlsx               : ``AFS_method_comparison_table.xlsx`` loaded as structured
    rows + the notes sheet, written to ``knowledge/method_table.json`` and tagged
    as S01 Table 1, page 6 (the transcription source).
  * ``knowledge/manifest.json`` records kb_version, sources, page counts, scopes.

Citation format is ``Sxx p.N``. Missing source files fail with a clear message.

PDF text extraction uses ``pypdf`` (falling back to ``pdfminer.six``); the xlsx
uses ``openpyxl`` with a dependency-free stdlib fallback. Install the pinned
dependencies from requirements.txt to run this end to end.

Usage:
    python ingest.py                 # build knowledge/ from materials/
    python ingest.py --check         # validate inputs without writing
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional
from xml.etree import ElementTree as ET

REPO_ROOT = Path(__file__).resolve().parent
MATERIALS_DIR = REPO_ROOT / "materials"
KNOWLEDGE_DIR = REPO_ROOT / "knowledge"

# Scope tags per source (mirrors spectro.kb.SOURCE_SCOPE / steering).
SOURCE_SCOPE = {
    "S01": "AFS technical",
    "S02": "AFS slides",
    "S03": "research paper",
    "S04": "research paper",
    "S05": "business context",
}

# Text PDFs to extract per page.
TEXT_PDFS = {
    "S01": "S01_Multispectral_Laboratory.pdf",
    "S03": "S03_LED_Kidney_Study_2017.pdf",
    "S04": "S04_Combi_MIR_Fluorescence_2021.pdf",
    "S05": "S05_AFS_Flyer.pdf",
}
XLSX_NAME = "AFS_method_comparison_table.xlsx"
SLIDES_NAME = "S02_slides.json"  # already committed in knowledge/

_SPREADSHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class IngestError(RuntimeError):
    """Raised with a clear message when ingestion cannot proceed."""


# ---------------------------------------------------------------------------
# PDF text extraction
# ---------------------------------------------------------------------------


def _extract_pdf_pages(path: Path) -> List[str]:
    """Return a list of per-page text strings (index 0 == page 1).

    Tries pypdf first, then pdfminer.six. Raises IngestError if neither is
    available or extraction fails.
    """
    # 1) pypdf
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "") for page in reader.pages]
        if any(p.strip() for p in pages):
            return pages
    except ImportError:
        pass
    except Exception as exc:  # pragma: no cover - depends on file/lib
        raise IngestError(f"pypdf failed to read {path.name}: {exc}") from exc

    # 2) pdfminer.six (per-page)
    try:
        from pdfminer.high_level import extract_pages  # type: ignore
        from pdfminer.layout import LTTextContainer  # type: ignore

        pages: List[str] = []
        for layout in extract_pages(str(path)):
            parts = [
                el.get_text() for el in layout if isinstance(el, LTTextContainer)
            ]
            pages.append("".join(parts))
        if pages:
            return pages
    except ImportError as exc:
        raise IngestError(
            f"Cannot extract text from {path.name}: install 'pypdf' or 'pdfminer.six' "
            f"(see requirements.txt)."
        ) from exc
    except Exception as exc:  # pragma: no cover
        raise IngestError(f"pdfminer failed to read {path.name}: {exc}") from exc

    raise IngestError(f"No extractable text found in {path.name}.")


def build_pdf_source(source_id: str, path: Path) -> dict:
    pages = _extract_pdf_pages(path)
    scope = SOURCE_SCOPE[source_id]
    entries = []
    for i, text in enumerate(pages, start=1):
        entries.append({"page": i, "text": _normalize_ws(text)})
    return {
        "source_id": source_id,
        "scope": scope,
        "source_file": path.name,
        "page_count": len(entries),
        "pages": entries,
    }


# ---------------------------------------------------------------------------
# xlsx extraction (openpyxl, with a stdlib fallback)
# ---------------------------------------------------------------------------


def _read_xlsx_openpyxl(path: Path) -> Dict[str, List[List[str]]]:
    from openpyxl import load_workbook  # type: ignore

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    out: Dict[str, List[List[str]]] = {}
    for ws in wb.worksheets:
        rows: List[List[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else str(c) for c in row])
        out[ws.title] = rows
    return out


def _read_xlsx_stdlib(path: Path) -> Dict[str, List[List[str]]]:
    """Dependency-free xlsx reader (handles shared strings + inline strings)."""
    z = zipfile.ZipFile(str(path))

    shared: List[str] = []
    if "xl/sharedStrings.xml" in z.namelist():
        root = ET.fromstring(z.read("xl/sharedStrings.xml"))
        for si in root.findall(_SPREADSHEET_NS + "si"):
            shared.append("".join(t.text or "" for t in si.iter(_SPREADSHEET_NS + "t")))

    # Map sheet name -> target file via workbook + rels.
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    rid_to_target = {
        r.get("Id"): r.get("Target")
        for r in rels.findall(
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
        )
    }
    names = set(z.namelist())

    def _resolve_target(target: str) -> str:
        """Resolve a worksheet rel Target to a real archive path.

        Targets may be absolute ("/xl/worksheets/sheet1.xml") or relative to the
        workbook's folder ("worksheets/sheet1.xml" -> "xl/worksheets/sheet1.xml").
        """
        t = target.lstrip("/")
        for candidate in (t, "xl/" + t, t.split("worksheets/")[-1]):
            if candidate in names:
                return candidate
            if ("xl/" + candidate) in names:
                return "xl/" + candidate
        return t

    name_to_file: Dict[str, str] = {}
    for sheet in wb.find(_SPREADSHEET_NS + "sheets").findall(_SPREADSHEET_NS + "sheet"):
        name = sheet.get("name")
        rid = sheet.get(rel_ns + "id")
        name_to_file[name] = _resolve_target(rid_to_target.get(rid, ""))

    out: Dict[str, List[List[str]]] = {}
    for name, fpath in name_to_file.items():
        root = ET.fromstring(z.read(fpath))
        rows: List[List[str]] = []
        for row in root.iter(_SPREADSHEET_NS + "row"):
            cells: List[str] = []
            for c in row.findall(_SPREADSHEET_NS + "c"):
                t = c.get("t")
                v = c.find(_SPREADSHEET_NS + "v")
                is_el = c.find(_SPREADSHEET_NS + "is")
                if t == "s" and v is not None:
                    val = shared[int(v.text)]
                elif is_el is not None:
                    val = "".join(x.text or "" for x in is_el.iter(_SPREADSHEET_NS + "t"))
                elif v is not None:
                    val = v.text or ""
                else:
                    val = ""
                cells.append(val)
            rows.append(cells)
        out[name] = rows
    return out


def build_method_table(path: Path) -> dict:
    try:
        sheets = _read_xlsx_openpyxl(path)
    except ImportError:
        sheets = _read_xlsx_stdlib(path)
    except Exception as exc:  # pragma: no cover
        raise IngestError(f"Failed to read {path.name}: {exc}") from exc

    # Identify the data sheet and the notes sheet by name, with fallbacks.
    data_name = next((n for n in sheets if "method" in n.lower() or "comparison" in n.lower()), None)
    notes_name = next((n for n in sheets if "note" in n.lower()), None)
    if data_name is None:
        # Fall back to the first sheet.
        data_name = list(sheets)[0]

    data_rows = sheets[data_name]
    if not data_rows:
        raise IngestError(f"{path.name}: sheet {data_name!r} has no rows.")
    header = [h.strip() for h in data_rows[0]]
    rows: List[Dict[str, str]] = []
    for raw in data_rows[1:]:
        if not any(str(c).strip() for c in raw):
            continue
        row = {header[i]: (raw[i] if i < len(raw) else "") for i in range(len(header))}
        rows.append(row)

    notes_lines: List[str] = []
    if notes_name:
        for raw in sheets[notes_name]:
            line = " ".join(str(c) for c in raw if str(c).strip())
            if line.strip():
                notes_lines.append(line.strip())

    flat_rows = " ; ".join(
        " ".join(f"{k}={v}" for k, v in row.items() if str(v).strip()) for row in rows
    )
    text = _normalize_ws(flat_rows + " " + " ".join(notes_lines))

    return {
        "source_id": "S01",
        "page": 6,
        "locator": "Table 1",
        "scope": SOURCE_SCOPE["S01"],
        "kind": "xlsx_rows",
        "source_file": path.name,
        "columns": header,
        "rows": rows,
        "notes": notes_lines,
        "text": text,
    }


# ---------------------------------------------------------------------------
# S02 slides
# ---------------------------------------------------------------------------


def load_slides(slides_path: Path) -> dict:
    """Validate and summarize the committed S02 slide notes (no vision)."""
    if not slides_path.exists():
        raise IngestError(
            f"Missing {slides_path} — the S02 slide notes must be committed (no vision pass)."
        )
    with open(slides_path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    slides = data.get("slides", [])
    if not slides:
        raise IngestError(f"{slides_path.name} has no 'slides' array.")
    numbers = sorted(int(s["slide_number"]) for s in slides)
    return {
        "source_id": "S02",
        "scope": SOURCE_SCOPE["S02"],
        "source_file": slides_path.name,
        "page_count": len(numbers),
        "slide_numbers": numbers,
        "provenance": data.get("provenance", ""),
    }


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

import re as _re

_WS_RE = _re.compile(r"\s+")


def _normalize_ws(text: str) -> str:
    return _WS_RE.sub(" ", text or "").strip()


def _check_inputs() -> List[str]:
    """Return a list of missing-input error messages (empty == all present)."""
    missing: List[str] = []
    for source_id, fname in TEXT_PDFS.items():
        if not (MATERIALS_DIR / fname).exists():
            missing.append(f"missing material for {source_id}: materials/{fname}")
    if not (MATERIALS_DIR / XLSX_NAME).exists():
        missing.append(f"missing material: materials/{XLSX_NAME}")
    if not (KNOWLEDGE_DIR / SLIDES_NAME).exists():
        missing.append(f"missing committed slide notes: knowledge/{SLIDES_NAME}")
    return missing


def _compute_kb_version(payloads: Dict[str, dict]) -> str:
    h = hashlib.sha256()
    for name in sorted(payloads):
        h.update(name.encode("utf-8"))
        h.update(json.dumps(payloads[name], sort_keys=True, ensure_ascii=False).encode("utf-8"))
    return "kb-" + h.hexdigest()[:12]


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def ingest(check_only: bool = False) -> dict:
    missing = _check_inputs()
    if missing:
        raise IngestError("Cannot ingest:\n  - " + "\n  - ".join(missing))

    if check_only:
        return {"status": "inputs_ok", "sources": sorted(list(TEXT_PDFS) + ["S02", "xlsx"])}

    KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)

    written: Dict[str, dict] = {}

    # Text PDFs.
    for source_id, fname in TEXT_PDFS.items():
        payload = build_pdf_source(source_id, MATERIALS_DIR / fname)
        out_path = KNOWLEDGE_DIR / f"{source_id}.json"
        _write_json(out_path, payload)
        written[out_path.name] = _manifest_entry(payload, out_path.name)

    # xlsx method table.
    table = build_method_table(MATERIALS_DIR / XLSX_NAME)
    table_path = KNOWLEDGE_DIR / "method_table.json"
    _write_json(table_path, table)
    written[table_path.name] = {
        "source_id": "S01",
        "scope": table["scope"],
        "kind": "xlsx_rows",
        "locator": "Table 1",
        "page": 6,
        "row_count": len(table["rows"]),
        "file": table_path.name,
    }

    # S02 slides (already committed; validate + record).
    slides_info = load_slides(KNOWLEDGE_DIR / SLIDES_NAME)
    written[SLIDES_NAME] = {
        "source_id": "S02",
        "scope": slides_info["scope"],
        "kind": "slide_notes",
        "page_count": slides_info["page_count"],
        "file": SLIDES_NAME,
    }

    manifest = {
        "kb_version": _compute_kb_version(
            {k: v for k, v in [(n, _read_json(KNOWLEDGE_DIR / n)) for n in written]}
        ),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "scopes": SOURCE_SCOPE,
        "sources": written,
    }
    _write_json(KNOWLEDGE_DIR / "manifest.json", manifest)
    return manifest


def _manifest_entry(payload: dict, filename: str) -> dict:
    return {
        "source_id": payload["source_id"],
        "scope": payload["scope"],
        "kind": "pdf_page",
        "page_count": payload.get("page_count", len(payload.get("pages", []))),
        "file": filename,
    }


def _write_json(path: Path, obj: dict) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False, indent=2)
        fh.write("\n")


def _read_json(path: Path) -> dict:
    with open(path, "r", encoding="utf-8") as fh:
        return json.load(fh)


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Build the page-tagged knowledge base.")
    parser.add_argument("--check", action="store_true", help="validate inputs without writing")
    args = parser.parse_args(argv)
    try:
        result = ingest(check_only=args.check)
    except IngestError as exc:
        print(f"ingest: ERROR\n{exc}", file=sys.stderr)
        return 2
    if args.check:
        print("ingest: inputs OK")
    else:
        print(f"ingest: wrote knowledge/ (kb_version={result['kb_version']})")
        for name, info in result["sources"].items():
            pc = info.get("page_count") or info.get("row_count")
            print(f"  - {name}: {info['source_id']} [{info['scope']}] ({pc})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
